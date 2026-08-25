import sys
import os
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from parse_curriculum import parse_curriculum_text
from build_full_dataset import CodeforcesMasterDataset, WORKBOOK_PATH

def run_kingdom_processor(input_source):
    """
    input_source can be a string containing the raw curriculum text or a path to a file.
    """
    if os.path.exists(input_source):
        with open(input_source, 'r', encoding='utf-8') as f:
            text = f.read()
    else:
        text = input_source

    curriculum_data = parse_curriculum_text(text)
    if not curriculum_data:
        print("No valid Kingdom / Pattern / Problem data found in input.")
        return

    builder = CodeforcesMasterDataset(WORKBOOK_PATH)

    # If existing workbook exists, load seen problems to preserve state and avoid duplication
    if os.path.exists(WORKBOOK_PATH):
        try:
            wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
            if "Problems" in wb.sheetnames:
                ws = wb["Problems"]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if row and len(row) >= 4 and row[3]:
                        pid = str(row[3]).strip().upper()
                        builder.seen_problems.add(pid)
                        # Also reconstruct problem rows
                        if len(row) >= 14:
                            builder.problems_data.append({
                                "Status": row[0] or "Unsolved",
                                "Category": row[1] or "",
                                "Subtopic": row[2] or "",
                                "Problem #": row[3] or "",
                                "Contest ID": row[4] or "",
                                "Problem Index": row[5] or "",
                                "Problem Name": row[6] or "",
                                "Rating": row[7] or "UNKNOWN",
                                "Difficulty": row[8] or "UNKNOWN",
                                "Topics": row[9] or "",
                                "Estimated Time (mins)": row[10] or 0,
                                "XP": row[11] or 0,
                                "Open Link": row[12] or "",
                                "Notes": row[13] or ""
                            })
                            kingdom = row[1]
                            subtopic = row[2]
                            if kingdom and kingdom not in builder.kingdom_order:
                                builder.kingdom_order.append(kingdom)
                            if kingdom not in builder.kingdom_patterns:
                                builder.kingdom_patterns[kingdom] = set()
                            if kingdom and subtopic:
                                builder.kingdom_patterns[kingdom].add(subtopic)
                                if (kingdom, subtopic) not in builder.pattern_order:
                                    builder.pattern_order.append((kingdom, subtopic))
            print(f"Loaded existing workbook with {len(builder.seen_problems)} unique problems.")
        except Exception as e:
            print(f"Note: Could not load existing workbook state: {e}")

    total_added = 0
    total_dups = 0
    total_unk = 0

    print("\n========================================================")
    print("      PROCESSING CODEFORCES CURRICULUM DATASET          ")
    print("========================================================\n")

    for k_item in curriculum_data:
        k_name = k_item["kingdom"]
        added, dups, unk, k_stats = builder.process_curriculum([k_item])
        
        total_added += added
        total_dups += dups
        total_unk += unk
        
        st = k_stats.get(k_name, {"added": added, "duplicates": dups, "unknown": unk})
        print(f"[OK] Kingdom Completed: {k_name}")
        print(f"  - Unique problems added: {st['added']}")
        print(f"  - Duplicates skipped:    {st['duplicates']}")
        print(f"  - Unknown problems:      {st['unknown']}")
        print(f"  - Current total rows:    {len(builder.problems_data)}")
        print("--------------------------------------------------------")

    print("\n========================================================")
    print(f"FINAL SUMMARY:")
    print(f"Total Unique Problems Added: {total_added}")
    print(f"Total Duplicates Skipped:   {total_dups}")
    print(f"Total Unknown Problems:     {total_unk}")
    print(f"Total Dataset Rows:         {len(builder.problems_data)}")
    print(f"Workbook Saved to:          {WORKBOOK_PATH}")
    print("========================================================\n")

if __name__ == "__main__":
    if len(sys.argv) > 1:
        run_kingdom_processor(sys.argv[1])
    else:
        print("Usage: python process_kingdoms.py <path_or_text>")
