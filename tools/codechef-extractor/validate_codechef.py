import json
import openpyxl
import os
import re

WORKBOOK_PATH = "CodeChef_Master_Dataset.xlsx"
SRC_JSON_PATH = "src/data/codechef.json"
PUBLIC_JSON_PATH = "public/data/codechef.json"

EXPECTED_TOPIC_TARGETS = {
    "Arrays": 40,
    "Prefix Sum": 30,
    "Two Pointers": 30,
    "Sliding Window": 30,
    "Binary Search": 40,
    "Sorting & Greedy": 40,
    "Strings": 35,
    "Bit Manipulation": 35,
    "Mathematics": 45,
    "Recursion": 25,
    "Stack": 35,
    "Queue": 20,
    "Linked List": 20,
    "Trees": 50,
    "DFS & BFS": 50,
    "Shortest Paths": 35,
    "DSU": 25,
    "Dynamic Programming": 70,
    "Divide & Conquer": 20,
    "Segment Tree": 35,
    "Fenwick Tree": 20,
    "Trie": 20,
    "Geometry": 20,
    "Game Theory": 15,
    "Advanced Topics": 30
}

def validate_all():
    print("=== STARTING CODECHEF DATASET VALIDATION ===")

    # 1. Check file existence
    for path in [WORKBOOK_PATH, SRC_JSON_PATH, PUBLIC_JSON_PATH]:
        if not os.path.exists(path):
            raise FileNotFoundError(f"Missing file: {path}")
        print(f"[OK] File exists: {path}")

    # 2. Validate JSON files
    with open(SRC_JSON_PATH, "r", encoding="utf-8") as f:
        src_data = json.load(f)

    with open(PUBLIC_JSON_PATH, "r", encoding="utf-8") as f:
        pub_data = json.load(f)

    assert len(src_data) == len(pub_data), "src/data/codechef.json and public/data/codechef.json counts differ!"
    print(f"[OK] Total problems in JSON: {len(src_data)}")

    # Check for duplicate problem codes
    seen_ids = set()
    url_pattern = re.compile(r"^https://www\.codechef\.com/problems/[A-Za-z0-9_]+$")
    difficulty_counts = {}
    kingdom_counts = {}
    pattern_counts = {}

    for idx, prob in enumerate(src_data, 1):
        pid = prob["id"]
        if pid in seen_ids:
            raise ValueError(f"Duplicate problem ID in JSON: {pid}")
        seen_ids.add(pid)

        # Validate URL
        url = prob["url"]
        if not url_pattern.match(url):
            raise ValueError(f"Invalid URL format for problem {pid}: {url}")

        # Validate Platform
        if prob["platform"] != "codechef":
            raise ValueError(f"Platform for {pid} is '{prob['platform']}' instead of 'codechef'")

        # Tally stats
        diff = prob["difficulty"]
        difficulty_counts[diff] = difficulty_counts.get(diff, 0) + 1

        k = prob["kingdom"]
        kingdom_counts[k] = kingdom_counts.get(k, 0) + 1

        p = prob["pattern"]
        pattern_counts[p] = pattern_counts.get(p, 0) + 1

    print("[OK] Zero duplicate problem codes.")
    print("[OK] 100% problem URLs match official CodeChef format.")
    print("[OK] Platform field is 100% 'codechef'.")

    # 3. Validate Excel Workbook
    wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=False)
    expected_sheets = ["Problems", "Kingdom Summary", "Pattern Summary", "Statistics", "Metadata"]
    for sname in expected_sheets:
        assert sname in wb.sheetnames, f"Sheet '{sname}' missing in workbook!"
    print("[OK] All 5 workbook sheets verified.")

    ws_problems = wb["Problems"]
    workbook_problem_count = ws_problems.max_row - 1
    assert workbook_problem_count == len(src_data), f"Workbook row count ({workbook_problem_count}) does not match JSON count ({len(src_data)})"
    print(f"[OK] Excel workbook contains exact {workbook_problem_count} problem rows.")

    # Print summary report
    print("\n==================================================")
    print("CODECHEF MASTER DATASET VERIFICATION SUMMARY")
    print("==================================================")
    print(f"Total Selected Problems: {len(src_data)}")
    print(f"Total Kingdoms Mapped:   {len(kingdom_counts)}")
    print(f"Total Patterns Mapped:   {len(pattern_counts)}")

    print("\n--- Difficulty Distribution ---")
    for d, count in sorted(difficulty_counts.items(), key=lambda x: x[1], reverse=True):
        print(f"  - {d:15s}: {count} ({count/len(src_data)*100:.1f}%)")

    print("\n--- Problems per Kingdom ---")
    for k, count in sorted(kingdom_counts.items()):
        print(f"  - {k}: {count}")

    print("\n[OK] ALL VALIDATION CHECKS PASSED SUCCESSFULLY!")

if __name__ == "__main__":
    validate_all()
